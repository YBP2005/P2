# -*- coding: utf-8 -*-
"""按 **PR 官方条款**实测页数（A 级条款，来自 `通用经验与教训\26_PR投稿指南…20260920.md` §3.1）：

    单栏（禁双栏）｜两端对齐｜页码
    正文与表 **10 pt** Times New Roman｜脚注与图注 8 pt
    行距：Word **1.5 倍**
    页边距：上 4.3 / 右 4.8 / 下 4.3 / 左 4.8 cm
    A4

此前所有页数都用 **12pt / 双倍 / 2.54cm**（26 号文 §5.3 自己写明"比要求更保守"）。
页数是外部硬约束，**唯一该拿去比 35 页上限的是官方格式下的数**，所以本脚本给出后者，
并同时给出此前口径，便于并列核对。只写 `work/_tmp_prgeo/`。
"""
import io
import os
import re
import sys

sys.stdout.reconfigure(encoding='utf-8')
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

W = r'E:\workplace'
TMP = os.path.join(W, 'work', '_tmp_prgeo')
os.makedirs(TMP, exist_ok=True)
SRC = os.path.join(W, 'P2_English_submission_blind_v1.md')

TOKEN = re.compile(r'(\*\*.+?\*\*|`[^`]+?`|\*[^*\n]+?\*)')
SPECIAL = re.compile(r'^(#|\||!|>|```|---|\*|- |\d+\. )')


def add_runs(p, text):
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            p.add_run(part[2:-2]).bold = True
        elif part.startswith('`') and part.endswith('`') and len(part) > 2:
            r = p.add_run(part[1:-1]); r.font.name = 'Consolas'
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            p.add_run(part[1:-1]).italic = True
        else:
            p.add_run(part)


def split_row(line):
    return [c.strip() for c in line.strip().strip('|').split('|')]


def build(out_path, body_pt=10.0, spacing=1.5, m=(4.3, 4.8, 4.3, 4.8)):
    md = io.open(SRC, encoding='utf-8').read().split('\n')
    doc = docx.Document()
    for sec in doc.sections:
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.top_margin, sec.right_margin = Cm(m[0]), Cm(m[1])
        sec.bottom_margin, sec.left_margin = Cm(m[2]), Cm(m[3])
    st = doc.styles['Normal']
    st.font.name = 'Times New Roman'
    st.font.size = Pt(body_pt)
    pf = st.paragraph_format
    pf.line_spacing = spacing
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY      # 官方：全篇两端对齐
    i = 0
    while i < len(md):
        t = md[i].strip()
        if t.startswith('```'):
            i += 1
            buf = []
            while i < len(md) and not md[i].strip().startswith('```'):
                buf.append(md[i]); i += 1
            i += 1
            add_runs(doc.add_paragraph(), '\n'.join(buf))
            continue
        if t.startswith('|') and i + 1 < len(md) and re.match(r'^\|[\s:\-|]+\|$', md[i + 1].strip()):
            rows = []
            while i < len(md) and md[i].strip().startswith('|'):
                rows.append(split_row(md[i])); i += 1
            rows = [r for r in rows if not re.match(r'^[\s:\-]+$', '|'.join(r))]
            if rows:
                tb = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
                for ri, r in enumerate(rows):
                    for ci, cell in enumerate(r):
                        if ci >= len(tb.rows[ri].cells):
                            continue
                        cp = tb.rows[ri].cells[ci].paragraphs[0]
                        rr = cp.add_run(cell); rr.font.size = Pt(body_pt)   # 官方：表 10pt
                        cp.paragraph_format.line_spacing = 1.0
            continue
        if t.startswith('![') and '](' in t:
            path = t[t.index('](') + 2:t.rindex(')')]
            full = os.path.join(W, path)
            if os.path.exists(full):
                doc.add_picture(full, width=Cm(11.4))     # 版心宽 21 − 4.8 − 4.8
            i += 1
            continue
        if t.startswith('#'):
            lvl = len(t) - len(t.lstrip('#'))
            h = t.lstrip('# ').strip().replace('**', '')
            p = doc.add_heading(h, level=min(lvl, 3))
            for r in p.runs:
                r.font.name = 'Times New Roman'
            i += 1
            continue
        if not t or t == '---':
            i += 1
            continue
        buf = [t]
        i += 1
        while i < len(md):
            nx = md[i].strip()
            if not nx or SPECIAL.match(nx) or nx.startswith('!['):
                break
            buf.append(nx)
            i += 1
        add_runs(doc.add_paragraph(), ' '.join(buf))
    for sec in doc.sections:
        run = sec.footer.paragraphs[0].add_run()
        run._r.append(run._r.makeelement(qn('w:fldSimple'), {qn('w:instr'): 'PAGE'}))
    doc.save(out_path)


def measure(app, path):
    doc = app.Documents.Open(path, False, True)
    doc.ExportAsFixedFormat(path[:-5] + '.pdf', 17)
    pages = doc.ComputeStatistics(2)
    words = doc.ComputeStatistics(0)
    doc.Close(False)
    return pages, words


def main():
    import win32com.client
    app = win32com.client.DispatchEx('Word.Application')
    try:
        app.Visible = False
        app.DisplayAlerts = 0
    except Exception:
        pass
    print()
    print('%-38s %6s %8s' % ('口径', '页数', 'Word 词数'))
    try:
        combos = [
            ('PR 官方（10pt/1.5倍/4.3-4.8cm）', 10.0, 1.5, (4.3, 4.8, 4.3, 4.8)),
            ('仅字号 12pt，其余照官方', 12.0, 1.5, (4.3, 4.8, 4.3, 4.8)),
            ('仅行距双倍，其余照官方', 10.0, 2.0, (4.3, 4.8, 4.3, 4.8)),
        ]
        for name, pt, sp, m in combos:
            out = os.path.join(TMP, 'pr_%gpt_%gsp.docx' % (pt, sp))
            build(out, pt, sp, m)
            p, w = measure(app, out)
            flag = 'IN ' if p <= 35 else 'OVER'
            print('%-38s %6d %8d   %s' % (name, p, w, flag))
    finally:
        try:
            app.Quit()
        except Exception:
            try:
                app.Application.Quit()
            except Exception:
                pass
    print()
    print('上限 35 页（含图、表、参考文献、bio-sketches；附录官方自相矛盾，取更严一侧）')
    print('对照：此前口径 12pt/双倍/2.54cm = 41 页（本目录另一脚本给出）')


if __name__ == '__main__':
    main()
